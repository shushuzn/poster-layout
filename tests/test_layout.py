# -*- coding: utf-8 -*-
"""版面不变量与产物的用例。

刻意**不依赖某一次调版式的具体数值**（字号、图宽这类），只断言必须成立的性质，
避免以后每改一次文案就要改测试。所有用例都跑在随仓库发布的**示例任务**
``tasks/demo`` 上（它由 ``tools/make_demo.py`` 生成，不含任何真实赛事材料）。

为什么显式写任务名：本机可能还放着别的（不入库的）任务，此时 ``find_task(None)``
会因为"tasks/ 下有多个任务"而报错；而克隆下来的人只有示例任务。写死 ``TASK``
两边都稳。
"""
from __future__ import annotations

import hashlib
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

from poster import config as C          # noqa: E402
from poster import layout as L          # noqa: E402
from poster import template as T        # noqa: E402
from poster.cli import prepare          # noqa: E402
from poster.imaging import prepare_figures  # noqa: E402

TASK = "demo"                           # 随仓库发布的示例任务
_CACHE: dict = {}


def _ctx():
    """任务 + 配图 + 版面，只算一次给所有用例复用。"""
    if not _CACHE:
        task = prepare(TASK)
        figs = prepare_figures(task)
        _CACHE["task"] = task
        _CACHE["figs"] = figs
        _CACHE["layout"] = L.build_layout(task, figs, verbose=False)
    return _CACHE["task"], _CACHE["figs"], _CACHE["layout"]


# ---------------------------------------------------------------- 任务与内容
def test_task_has_sections_and_figures():
    task, figs, _ = _ctx()
    assert len(task.sections) >= 1
    assert len(figs) == len(task.sections)
    assert task.char_count() > 0


def test_grid_covers_every_section():
    task, _, _ = _ctx()
    flat = [r for row in task.rows for r in row]
    assert sorted(flat) == list(range(len(task.sections)))


def test_no_missing_glyphs():
    """正文不得出现模板字体缺失的字形（历史踩过 ℃ / ₓ / ⁺）。"""
    task, _, _ = _ctx()
    bad = set("℃ₓ⁺⁻²³")
    text = "".join(t for sec in task.sections for para in sec.paras for t, _, _ in para)
    text += "".join(task.titles) + "".join(c for caps in task.captions for c in caps)
    hit = sorted(set(text) & bad)
    assert not hit, f"含字体缺失字形: {hit}"


# ---------------------------------------------------------------- 版面
def test_layout_is_feasible():
    """总高必须装进内容区，不允许溢出。"""
    _, _, lay = _ctx()
    assert lay.total <= C.AVAIL_H + 1e-6, f"溢出 {lay.total - C.AVAIL_H:.3f}cm"
    assert C.BODY_MIN - 1e-9 <= lay.body_pt <= C.BODY_MAX + 1e-9


def test_blocks_stay_inside_content_box():
    """每个块都必须落在内容区内。"""
    _, _, lay = _ctx()
    for c in lay.cells:
        assert c.x >= C.X0 - 1e-6, f"#{c.r} 左越界"
        assert c.x + c.bw <= C.X1 + 1e-6, f"#{c.r} 右越界"
        assert c.y >= C.Y0 - 1e-6, f"#{c.r} 上越界"
        assert c.y + c.h <= C.Y1 + 1e-6, f"#{c.r} 下越界"


def test_text_column_not_crushed():
    """正文列宽必须有下限，否则会挤成细长条（历史问题）。"""
    _, _, lay = _ctx()
    for c in lay.cells:
        assert c.tw >= C.TW_MIN - 1e-6, f"#{c.r} 正文列仅 {c.tw:.2f}cm"


def test_layout_mode_matches_figure_shape():
    """没写死排法的块，排法必须由图条的宽高比决定，不能一刀切。"""
    task, figs, lay = _ctx()
    by_r = {c.r: c for c in lay.cells}
    for r, (sec, bf) in enumerate(zip(task.sections, figs)):
        ratio = L.strip_ratio(bf)
        if ratio is None:
            assert by_r[r].style == "text", f"#{r} 无配图应为纯文本块"
        elif sec.mode != "auto":
            assert by_r[r].style == sec.mode, f"#{r} 应服从文档里写死的 {sec.mode}"
        else:
            w_over_h = 1 / ratio
            if w_over_h > 3.0:                  # 超宽条 → 必须通栏横幅
                assert by_r[r].style == "stack", f"#{r} 超宽图应为 stack"
            elif w_over_h <= 1.20:              # 方图 → 必须左文右图
                assert by_r[r].style == "side", f"#{r} 方图应为 side"


def test_figure_strip_fits_block_width():
    """图条（含张间间距）不得超出块内可用宽，否则会压到邻块。"""
    _, figs, lay = _ctx()
    for c in lay.cells:
        if not c.has_fig:
            continue
        assert c.fw <= c.inner + 1e-6, f"#{c.r} 图条 {c.fw:.2f} 超过块内宽 {c.inner:.2f}"
        w = sum(figs[c.r].widths(c.ih)) + max(0, len(figs[c.r].paths) - 1) * C.FIG_GAP_IN
        assert abs(w - c.fw) < 1e-6, f"#{c.r} 图条宽度与各图之和 {w:.3f} 不一致"


def test_no_figure_block_has_zero_image():
    """纯文本块不能残留图片占位。"""
    _, _, lay = _ctx()
    for c in lay.cells:
        if c.style == "text":
            assert c.fw == 0.0 and c.ih == 0.0


def test_area_auto_detection_is_sane():
    """auto 探测出的内容区必须落在画布内且尺寸合理。"""
    task, _, _ = _ctx()
    w, h = C.CANVAS_CM
    assert 0 < C.X0 < C.X1 <= w + 1e-6
    assert 0 < C.Y0 < C.Y1 <= h + 1e-6
    assert C.W > 3 and C.H > 5, "内容区太小，八成是占位框探测错了"
    assert T.detect_area(task.template) is not None


# ---------------------------------------------------------------- 产物
_RESULT: dict = {}


def _build():
    if not _RESULT:
        from poster.cli import build
        _RESULT.update(build(TASK, with_preview=False, verbose=False))
    return _RESULT["pptx"]


def test_pptx_generation_and_template_fields_untouched():
    """构建产物形状数正确，且模板原有文本一字未被改。

    做法：拿模板自身的非空文本做基准集合（**排除**会被主动清空的占位框），
    断言产物里全部仍在 —— 比硬编码字段名更稳。
    """
    from pptx import Presentation

    from poster.writer import is_placeholder

    def texts(path, drop_placeholder=False):
        prs = Presentation(str(path))
        out = []
        for sh in prs.slides[0].shapes:
            if not (sh.has_text_frame and sh.text_frame.text.strip()):
                continue
            if drop_placeholder and is_placeholder(sh):
                continue
            out.append(sh.text_frame.text.strip())
        return out

    task, _, _ = _ctx()
    template_texts = set(texts(task.template, drop_placeholder=True))
    joined = "\n".join(texts(_build()))

    missing = sorted(t for t in template_texts if t not in joined)
    assert not missing, f"模板原有文本被改动/删除: {missing}"

    for t in task.titles:
        assert t in joined, f"缺少小标题「{t}」"
    for caps in task.captions:                 # 一张图一条图注，逐条都要在产物里
        for c in caps:
            if c:
                assert c in joined, f"缺少图注「{c}」"


def test_pptx_has_figures_when_task_declares_them():
    """文档里声明了几张图，产物里就必须有几张 —— 多一块少一张都算错配。"""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    task, _, _ = _ctx()
    want = sum(len(s.figure_paths) for s in task.sections)
    prs = Presentation(str(_build()))
    got = sum(1 for sh in prs.slides[0].shapes
              if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    # 模板自带 1 张背景图，其余是排进去的配图
    assert got == want + 1, f"配图数量不对：期望 {want + 1}，实得 {got}"


def test_every_figure_sits_in_its_own_block():
    """图文对位：每张图都必须落在声明它的那个块的矩形内。

    这是本期返工的起因（一张表征图曾被排到不是讨论它的那个块下），
    用几何包含关系把它钉死，而不是靠肉眼看预览图。
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    task, figs, lay = _ctx()
    prs = Presentation(str(_build()))

    placed = [(round(sh.left / 360000, 3), round(sh.top / 360000, 3),
               round(sh.width / 360000, 3), round(sh.height / 360000, 3))
              for sh in prs.slides[0].shapes
              if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    placed = [p for p in placed if p[2] < 7.9]          # 排除模板整页背景图
    assert len(placed) == sum(len(f.paths) for f in figs), \
        f"排进去的图有 {len(placed)} 张，文档声明了 {sum(len(f.paths) for f in figs)} 张"

    for c in lay.cells:
        bf = figs[c.r]
        inside = [t for t in placed
                  if c.x - 0.01 <= t[0] <= c.x + c.bw + 0.01
                  and c.y - 0.01 <= t[1] <= c.y + c.h + 0.01]
        assert len(inside) == len(bf.paths), \
            "#%d「%s」块内实得 %d 张图，文档声明 %d 张" % (
                c.r, task.titles[c.r], len(inside), len(bf.paths))
        for x, y, w, h in inside:
            assert x >= c.x + C.BP_X - 1e-6, f"#{c.r} 图条左溢出块内边距"
            assert x + w <= c.x + c.bw - C.BP_X + 1e-6, f"#{c.r} 图条右溢出块内边距"
            assert abs(h - c.ih) < 5e-3, f"#{c.r} 图高 {h:.3f} 与求解值 {c.ih:.3f} 不符"


def test_build_is_reproducible():
    """同输入必须产出同字节。

    否则每跑一次构建就会在 git 里多留一份 ~16MB 的 pptx 历史
    —— 这正是本仓库 .git 曾膨胀到 265MB 的主因（见 repro.freeze_zip）。
    """
    from poster.cli import build

    path = _build()
    first = hashlib.sha256(path.read_bytes()).hexdigest()
    build(TASK, with_preview=False, verbose=False)     # 再构建一次
    assert hashlib.sha256(path.read_bytes()).hexdigest() == first, \
        "构建不可复现：同输入产出了不同字节"


# ---------------------------------------------------------------- 渲染器一致性
# 下面三条守的是"模型与渲染器对不齐"这一类错 —— 2026-10-14 的版面事故
# （末行被配图盖住）就是模型自洽、渲染器不认账。详见 docs/rendering-notes.md。
def test_punctuation_policy_matches_renderer():
    """默认 push 策略：行尾放不下的标点**推到下一行**（PowerPoint 的做法）。

    hang 策略（挤在行尾）只有在 pptx 里显式开了 ``hangingPunct`` 时才成立，
    所以行数估算绝不能用它 —— 少算一行，末行就会落到图底下。
    """
    from poster import typography as TY

    txt = "一二三四五六七八九十、"          # 10 个全角字正好占满 1.20cm，标点放不下
    hang = TY.wrap_para([(txt, "x")], 3.45, 1.20, "hang")
    push = TY.wrap_para([(txt, "x")], 3.45, 1.20, "push")
    assert len(hang) == 1 and hang[0][0][-1][0] == "、"
    assert len(push) == 2 and push[1][0][0][0] == "、"


def test_worst_policy_never_underestimates():
    """求解用的行数取两种标点策略里更悲观的那个。"""
    from poster import typography as TY

    txt = "一二三四五六七八九十、一二三四五六七八九十、"
    n = TY.lines_count_worst([[(txt, "x")]], 3.45, 1.20)
    assert n >= TY.lines_count([[(txt, "x")]], 3.45, 1.20, "hang")
    assert n >= TY.lines_count([[(txt, "x")]], 3.45, 1.20, "push")


def test_line_height_equals_written_spc_pct():
    """写进 pptx 的行距（spcPct）换算回来必须等于引擎的行高模型。

    这条是行距写法的护栏：一旦有人图省事改回 ``spcPts``，PowerPoint 会把
    5.52pt 四舍五入成 6.0pt，两边的行距就对不上了。
    """
    from poster import typography as TY

    for size in (2.6, 3.3, 3.45, 4.2):
        pct = TY.line_spacing_pct(C.LINE_FACTOR)
        pitch = size * C.PT2CM * C.LINE_REF_FACTOR * pct / 100000.0
        assert abs(pitch - TY.line_height(size)) < 1e-6, (
            "字号 %.2f：spcPct 换算行距 %.5f ≠ 模型行高 %.5f"
            % (size, pitch, TY.line_height(size)))


def test_script_runs_are_written_as_baseline():
    """化学式角标必须真的落成 ``a:rPr/@baseline``，不能只是解析出来。

    角标只位移、不缩放（字号不变），所以求解器的断行估算无需区分角标 ——
    但"模型算对了"不等于"渲染器排对了"，这条从**产物 XML** 上验。
    """
    import xml.etree.ElementTree as ET

    from poster import writer as W

    w = W.PosterWriter(C.template_path())
    w.add_text(1.0, 1.0, 4.0, 1.0,
               [[("Fe", C.BLACK, ""), ("3+", C.BLACK, "sup"),
                 (" 与 m", C.BLACK, ""), ("2", C.BLACK, "sub")]],
               3.0)
    xml = w.prs.slides[0].shapes[-1].text_frame._txBody.xml
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    runs = ET.fromstring(xml).findall(".//%sr" % ns)
    base = [r.find("%srPr" % ns).get("baseline") for r in runs]
    got = [r.find("%st" % ns).text for r in runs]
    assert got == ["Fe", "3+", " 与 m", "2"], got
    assert base[0] is None and base[2] is None, "普通文字不该带 baseline"
    assert base[1] == "30000" and base[3] == "-25000", base


def test_script_offset_is_one_glyph_scale():
    """预览里的角标位移必须落在字高量级内（上限 = 半个字高）。

    回归用例：``_script_dy`` 曾把 ``PX_PER_CM`` 与 ``scale`` 相乘两次，
    3.65pt 的上标被算成 -1545px、直接画到画布外 —— 预览里 ``Fe`` 后只剩一块
    空白，看着像"角标没渲染"，而产物其实是好的。这条锁住"预览与实排同口径"。
    """
    from poster import preview as P

    sc = C.PX_PER_CM
    glyph_px = 3.65 * C.PT2CM * sc            # 3.65pt 字高（px）
    sup = P._script_dy("sup", 3.65, sc)
    sub = P._script_dy("sub", 3.65, sc)
    assert sup < 0 < sub, (sup, sub)          # 上标向上、下标向下
    assert abs(sup) < glyph_px / 2 and abs(sub) < glyph_px / 2, (sup, sub, glyph_px)
    assert P._script_dy("", 3.65, sc) == 0.0
    # 位移与字号成正比：换字号只改幅度，不改口径
    assert abs(P._script_dy("sup", 7.3, sc)) > abs(sup), "位移应随字号线性变化"
