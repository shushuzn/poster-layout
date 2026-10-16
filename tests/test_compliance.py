# -*- coding: utf-8 -*-
"""合规清单的用例：清单不许与代码脱节，硬性项必须全绿。

设计要点
--------
* **文档与代码互为镜像**：`docs/compliance.md` 里出现的编号，与
  `compliance.REQUIREMENTS` 里的编号，必须完全一致 —— 加了检查不写文档、
  或写了文档不实现，这里就红。这是"清单不腐烂"的唯一保证。
* **硬性项全绿**：本任务在随仓库发布的状态下，32 条硬性要求必须全部通过
  （含 C8~C11 的"渲染器一致性"与 F5 的真渲染器复核；没装 PowerPoint 时 F5 跳过）。
* 不把字号 / 图高这类数值写死，只断言"要求成立"，免得每改一次文案就要改测试。
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

from poster import compliance as K          # noqa: E402
from poster.cli import build                # noqa: E402

_DOC = ROOT / "docs" / "compliance.md"
TASK = "demo"                           # 随仓库发布的示例任务
_CACHE: dict = {}


def _ctx():
    """构建一次，自检一次（跳过会再构建一遍的 F3，它由下面的用例单独守）。"""
    if not _CACHE:
        r = build(TASK, with_preview=False, verbose=False)
        ctx = K.load(r["task_obj"], r["figs"], r["layout"], r["pptx"])
        _CACHE["ctx"] = ctx
        _CACHE["results"] = K.run(ctx, skip=("F3",))
    return _CACHE["ctx"], _CACHE["results"]


def _doc_ids() -> list[str]:
    """取出 docs/compliance.md **要求表**里出现的全部编号。

    只扫到「人工项判读依据」之前 —— 那一节会重复列出人工项的编号（是判据说明，
    不是第二份要求），若一并统计会被当成重复编号。
    """
    txt = _DOC.read_text(encoding="utf-8").split("## 人工项判读依据")[0]
    return re.findall(r"^\|\s*([A-G]\d+)\s*\|", txt, re.M)


# ---------------------------------------------------------------- 清单本身
def test_doc_lists_match_code():
    """文档里的编号 与 代码里的编号 必须完全一致（互为镜像，不许脱节）。"""
    doc, code = set(_doc_ids()), set(K.BY_ID)
    assert doc == code, (
        "文档与代码不一致 —— 只在代码里: %s；只在文档里: %s"
        % (sorted(code - doc), sorted(doc - code)))


def test_doc_covers_every_id():
    """每条要求都要在文档里有独立表格行（用计数发现重复漏写）。"""
    ids = _doc_ids()
    assert len(ids) == len(set(ids)), "文档里有重复编号: %s" % ids
    assert len(ids) == len(K.BY_ID)


def test_items_are_well_formed():
    """编号唯一、分级合法、分组已知，且 auto 项真的挂了检查函数。"""
    seen = set()
    for it in K.REQUIREMENTS:
        assert it.id not in seen, "编号重复: %s" % it.id
        seen.add(it.id)
        assert it.level in ("hard", "soft", "manual"), it.id
        assert it.group in K.GROUPS, "%s 的分组 %r 未知" % (it.id, it.group)
        assert it.source, "%s 没写出处（模板/样张/评审/工程）" % it.id
        if it.level == "manual":
            assert it.func is None, "%s 是人工项，不该挂检查函数" % it.id
        else:
            assert callable(it.func), "%s 缺检查函数" % it.id


def test_manual_items_are_documented():
    """人工项必须在文档「人工项判读依据」里有判据，否则等于没有。"""
    txt = _DOC.read_text(encoding="utf-8")
    tail = txt.split("## 人工项判读依据")[-1]
    for it in K.REQUIREMENTS:
        if it.level == "manual":
            assert it.id in tail, "人工项 %s 没写判读依据" % it.id


def test_every_check_runs_without_error():
    """每条检查都要能跑通 —— 校验器自己崩了不该算「通过」。"""
    _, results = _ctx()
    crashed = [r.item.id for r in results if r.detail.startswith("校验自身出错")]
    assert not crashed, "检查函数抛异常: %s" % crashed


# ---------------------------------------------------------------- 结论
def test_all_hard_requirements_pass():
    """随仓库发布的这个任务，硬性要求必须全绿。"""
    _, results = _ctx()
    bad = K.failures(results)
    assert not bad, "\n".join("  ❌ %s %s — %s" % (r.item.id, r.item.title, r.detail)
                              for r in bad)


def test_soft_requirements_are_reported_separately():
    """提示项失败不算不合格，但必须出现在报告里（不许被静默吞掉）。"""
    _, results = _ctx()
    soft = [r for r in results if r.item.level == "soft"]
    assert soft, "清单里应有提示项"
    assert all(r.ok is not None for r in soft)


def test_report_covers_every_requirement():
    """报告里每条要求都要有编号与结论行（构建日志可当证据留存）。"""
    ctx, results = _ctx()
    text = K.report(results, ctx.task, ctx)
    for r in results:
        assert "| %s |" % r.item.id in text, "报告缺少 %s" % r.item.id
    assert "合规自检报告" in text


def test_build_is_reproducible_f3():
    """F3 单独跑一次（会再构建一遍，所以没放进 _ctx 的批量自检里）。"""
    ctx, _ = _ctx()
    ok, detail = K.BY_ID["F3"].func(ctx)
    assert ok, detail


# ---------------------------------------------------------------- 守卫能不能真的红
# 只断言"现在全绿"是不够的 —— 检查函数必须有**失败的能力**，否则它就是个摆设。
# 下面每条都把产物改坏一点点，证明对应的守卫会亮红灯。
def _mutate(ctx, fn) -> Path:
    """复制产物到临时文件、按 ``fn(prs, n_tpl)`` 改坏、返回新路径。

    ``n_tpl`` = 模板原有形状数：产物里前 ``n_tpl`` 个原封不动来自模板，
    改名换姓的破坏要**只落在引擎排进去的形状上**。它取自 ``ctx.n_tpl``，
    与 compliance 判定用的口径完全一致（不要写死 6 —— 模板是可换的）。
    """
    import shutil
    import tempfile

    from pptx import Presentation

    tmp = Path(tempfile.mkdtemp()) / ctx.out.name
    shutil.copy(ctx.out, tmp)
    prs = Presentation(str(tmp))
    fn(prs, ctx.n_tpl)
    prs.save(str(tmp))
    return tmp


def _check_against(ctx, mutated: Path, item_id: str):
    """拿改坏后的产物重跑某一条检查。"""
    from poster import compliance as K2

    c2 = K2.load(ctx.task, ctx.figs, ctx.lay, mutated)
    return K2.BY_ID[item_id].func(c2)


def _strip_lnspc_ours(prs, n_tpl):
    """把本引擎加的文本框的行距从 spcPct 换成 spcPts（历史事故的写法）。"""
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import qn

    ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    for sh in list(prs.slides[0].shapes)[n_tpl:]:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            pPr = p._p.find(qn("a:pPr"))
            if pPr is None:
                continue
            ln = pPr.find(qn("a:lnSpc"))
            if ln is None:
                continue
            pPr.remove(ln)
            pPr.insert(0, parse_xml('<a:lnSpc %s><a:spcPts val="552"/></a:lnSpc>' % ns))


def test_c9_flags_spcpts():
    """C9 必须能发现 spcPts —— 它会被 PowerPoint 四舍五入到整点，是上次翻车的根因。"""
    ctx, _ = _ctx()
    bad = _check_against(ctx, _mutate(ctx, _strip_lnspc_ours), "C9")
    assert bad[0] is False, "换成 spcPts 后 C9 竟然还通过"
    assert "spcPts" in bad[1]


def test_c8_flags_short_box():
    """C8 必须能发现「声明框高不够」—— 渲染器会把多出来的那行画到框外（压到图上）。"""
    def shrink(prs, n_tpl):
        # 取引擎排的**文字最多的那个框**（就是正文框）砍一半。
        # 不要按"超过 200 字"这种魔数找，换份文案就失效了。
        body = [s for s in list(prs.slides[0].shapes)[n_tpl:]
                if s.has_text_frame and s.text_frame.text.strip()]
        assert body, "产物里找不到引擎排进去的文本框，用例本身失效了"
        s = max(body, key=lambda x: len(x.text_frame.text))
        s.height = int(s.height * 0.5)

    ctx, _ = _ctx()
    bad = _check_against(ctx, _mutate(ctx, shrink), "C8")
    assert bad[0] is False, "把正文框高砍一半后 C8 竟然还通过"


def test_c10_flags_double_autofit():
    """C10 必须能发现 noAutofit 与 spAutoFit 并存 —— 那会让 PowerPoint 拒绝打开整个文件。"""
    def add_autofit(prs, n_tpl):
        from pptx.oxml import parse_xml
        from pptx.oxml.ns import qn

        ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        for s in list(prs.slides[0].shapes)[n_tpl:]:
            if not s.has_text_frame:
                continue
            bp = s.text_frame._txBody.find(qn("a:bodyPr"))
            if bp is not None and bp.find(qn("a:noAutofit")) is not None:
                bp.append(parse_xml('<a:spAutoFit %s/>' % ns))   # 与 noAutofit 并存
                return
        raise AssertionError("没找到挂了 noAutofit 的形状，用例本身失效了")

    ctx, _ = _ctx()
    bad = _check_against(ctx, _mutate(ctx, add_autofit), "C10")
    assert bad[0] is False, "挂了两个自动调整子元素，C10 竟然还通过"


def test_c11_flags_missing_hanging_punct():
    """C11 必须能发现悬挂标点没开 —— 那会让行首冒出「，。、」。"""
    def strip_attr(prs, n_tpl):
        from pptx.oxml.ns import qn

        for s in list(prs.slides[0].shapes)[n_tpl:]:
            if s.has_text_frame:
                bp = s.text_frame._txBody.find(qn("a:bodyPr"))
                if bp is not None and bp.get("hangingPunct"):
                    del bp.attrib["hangingPunct"]
                    return

    ctx, _ = _ctx()
    bad = _check_against(ctx, _mutate(ctx, strip_attr), "C11")
    assert bad[0] is False, "去掉 hangingPunct 后 C11 竟然还通过"


def test_f5_flags_text_under_picture():
    """F5（真渲染器复核）必须能发现「文字被图盖住」。

    做法就是把一张配图往上挪，压到它上面那段正文的末行上 —— 正是上次的事故形态。
    """
    from poster import renderer_check as RC

    if not RC.available():
        print("        （本机没有 PowerPoint，跳过）")
        return

    def lift_pic(prs, _n_tpl):
        for s in prs.slides[0].shapes:
            if s.shape_type == 13 and s.width / 360000 < 3.0:     # 各块配图（非整页底图）
                s.top = s.top - int(0.30 * 360000)

    ctx, _ = _ctx()
    ctx2 = K.load(ctx.task, ctx.figs, ctx.lay, _mutate(ctx, lift_pic))
    ok, detail = K.BY_ID["F5"].func(ctx2)
    assert ok is False, "把图往上挪 0.3cm 压住文字后，F5 竟然还通过"
    assert "压到图" in detail


def test_e6_flags_text_overload():
    """E6 必须能发现「字数超限」—— 字越多，字号被压得越小，上了墙谁也看不清。

    上限不是拍的：示例（``assets/ref`` 的样张）全篇约 250~300 汉字、
    最重的文本块约 130 字，这就是 ``TEXT_CHAR_BLOCK_MAX / TOTAL_MAX`` 的来处。
    """
    import copy

    from poster import compliance as K2

    ctx, _ = _ctx()
    fat = copy.deepcopy(ctx.task)
    for s in fat.sections:
        s.paras[0].append(("超" * (K2.C.TEXT_CHAR_BLOCK_MAX + 1), "black", ""))

    ok, detail = K2.BY_ID["E6"].func(K2.load(fat, ctx.figs, ctx.lay, ctx.out))
    assert ok is False, "把每个分块都塞到超过单块上限后，E6 竟然还通过"
    assert "超限" in detail


def test_e7_flags_ascii_chemistry():
    """E7 必须能发现「化学式没写成角标」—— ``Fe3+`` 上墙即排版错误。

    角标只位移、不缩放，所以这条不影响求解器的断行估算；
    只拦确定是排版错误的形态（离子电荷、氧化物、m2、ID/IG），不误伤普通文本。
    """
    import copy

    from poster import compliance as K2

    ctx, _ = _ctx()
    bad = copy.deepcopy(ctx.task)
    for s in bad.sections:
        s.paras[0].append(("Fe3+ 与 12.34 元/m2", "black", ""))
    ok, detail = K2.BY_ID["E7"].func(K2.load(bad, ctx.figs, ctx.lay, ctx.out))
    assert ok is False, "正文里写了 ASCII 化学式，E7 竟然还通过"
    assert "Fe3+" in detail and "m2" in detail, detail

    good = copy.deepcopy(ctx.task)
    for s in good.sections:
        s.paras[0].append(("Fe", "black", ""))
        s.paras[0].append(("3+", "black", "sup"))
        s.paras[0].append((" 与 12.34 元/m", "black", ""))
        s.paras[0].append(("2", "black", "sup"))
    ok2, _ = K2.BY_ID["E7"].func(K2.load(good, ctx.figs, ctx.lay, ctx.out))
    assert ok2 is True, "写成角标后 E7 不该报错"

