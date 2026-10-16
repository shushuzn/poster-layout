# -*- coding: utf-8 -*-
"""任务文档解析的用例。

这里刻意**不依赖随仓库发布的那份任务文档的具体文案** —— 除了一个"真任务能解析"
的连通性用例，其余全部用临时合成的文档，保证引擎换了墙报也照样成立。
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

from poster import config as C          # noqa: E402
from poster import spec                 # noqa: E402

TASK = "demo"                                     # 随仓库发布的示例任务
# 合成用例要一份**真实存在且随仓库发布**的模板：显式指向示例任务自带的模板
# （不要用 C.template_path()，它扫的是 assets/template，而赛事模板不入库）。
TEMPLATE = C.TASKS / TASK / "assets" / "template.pptx"


def _write(tmp: str, text: str, name: str = "task.md") -> Path:
    p = Path(tmp) / name
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(text, encoding="utf-8")
    return p


def _minimal_doc(body: str, extra: str = "") -> str:
    return (
        "---\n"
        "id: tmp-task\n"
        "title: 临时任务\n"
        'template: "%s"\n'
        "output: 临时墙报.pptx\n"
        "dist: out\n"
        "%s"
        "---\n\n%s" % (TEMPLATE.as_posix(), extra, body)
    )


# ---------------------------------------------------------------- 富文本
def test_emphasis_becomes_red():
    segs = spec.rich_text("甲**乙**丙")
    assert segs == [("甲", "black", ""), ("乙", "red", ""), ("丙", "black", "")], segs


def test_emphasis_at_edges_and_multiple():
    segs = spec.rich_text("**起**中**结**")
    assert [c for _, c, _ in segs] == ["red", "black", "red"], segs
    assert "".join(t for t, _, _ in segs) == "起中结"


def test_superscript_and_subscript_markup():
    """化学式/单位必须能写成真角标：^{} 上标、_{} 下标。"""
    segs = spec.rich_text("样品浸渍 Fe^{3+} 后")
    assert "".join(t for t, _, _ in segs) == "样品浸渍 Fe3+ 后"
    assert [s for _, _, s in segs] == ["", "sup", ""], segs
    segs = spec.rich_text("Raman I_{D}/I_{G} 降至 0.42")
    assert [s for _, _, s in segs if s] == ["sub", "sub"], segs
    assert "".join(t for t, _, _ in segs) == "Raman ID/IG 降至 0.42"


def test_script_markup_can_be_escaped():
    """``\\^{`` 输出字面量，不当标记解析。"""
    segs = spec.rich_text(r"字面量 \^{x}")
    assert [s for _, _, s in segs if s] == [], segs
    assert "^{x}" in "".join(t for t, _, _ in segs)


def test_plain_text_kept_verbatim():
    s = "含 数字 42 与符号 %/°C（括号）——破折号。"
    assert spec.rich_text(s) == [(s, "black", "")]


def test_rich_returns_rgb_after_config():
    """颜色标签要在用的时候才换成 RGB，这样任务覆盖配色也能生效。"""
    sec = spec.Section(title="x", paras=[[("红", "red", ""), ("黑", "black", "")]])
    assert sec.rich()[0][0][1] == C.RED
    assert sec.rich()[0][1][1] == C.BLACK


# ---------------------------------------------------------------- 正文与小节
def test_sections_titles_and_paragraphs():
    doc = _minimal_doc(
        "## 甲\n第一段。\n\n第二段。\n\n## 乙\n只有一段。\n")
    with tempfile.TemporaryDirectory() as tmp:
        t = spec.parse_task(_write(tmp, doc))
    assert t.titles == ["甲", "乙"]
    assert len(t.sections[0].paras) == 2, t.sections[0].paras
    assert t.sections[1].text() == "只有一段。"


def test_soft_wrap_inside_paragraph_is_joined():
    """段落内换行无意义，直接拼起来（中文排版）。"""
    doc = _minimal_doc("## 甲\n上半句\n下半句\n")
    with tempfile.TemporaryDirectory() as tmp:
        t = spec.parse_task(_write(tmp, doc))
    assert t.sections[0].text() == "上半句下半句"


def test_multiline_html_comment_is_stripped():
    doc = _minimal_doc("## 甲\n正文。\n<!--\n这里是跨行注释\nfigure: 不该被当成指令\n-->\n")
    with tempfile.TemporaryDirectory() as tmp:
        t = spec.parse_task(_write(tmp, doc))
    assert t.sections[0].text() == "正文。"
    assert t.sections[0].figures == []


def test_empty_section_rejected():
    doc = _minimal_doc("## 甲\n<!-- figure: none -->\n")
    with tempfile.TemporaryDirectory() as tmp:
        try:
            spec.parse_task(_write(tmp, doc))
        except ValueError as e:
            assert "没有正文" in str(e)
        else:
            raise AssertionError("空小节应当报错")


# ---------------------------------------------------------------- 配图指令
def _with_fig(tmp: str, directive: str, name: str = "img.png"):
    (Path(tmp) / "assets").mkdir(parents=True, exist_ok=True)
    (Path(tmp) / "assets" / name).write_bytes(b"\x89PNG\r\n\x1a\n")
    doc = _minimal_doc("## 甲\n%s\n正文。\n" % directive)
    return spec.parse_task(_write(tmp, doc))


def test_figure_directive_with_caption():
    with tempfile.TemporaryDirectory() as tmp:
        t = _with_fig(tmp, "<!-- figure: img.png | 图1  说明 -->")
        assert t.sections[0].figures == ["img.png"]
        assert t.sections[0].captions == ["图1  说明"]
        assert t.sections[0].figure_paths[0].is_file()
        assert t.sections[0].mode == "auto"


def test_multiple_figures_in_one_block():
    """单块多图：**分号**分隔，**每张图各写一条图注**，顺序即并排顺序。"""
    with tempfile.TemporaryDirectory() as tmp:
        (Path(tmp) / "assets").mkdir(parents=True, exist_ok=True)
        for n in ("a.png", "b.png", "c.png"):
            (Path(tmp) / "assets" / n).write_bytes(b"\x89PNG\r\n\x1a\n")
        doc = _minimal_doc("## 甲\n"
                           "<!-- figure: a.png | 图1 甲 ; b.png | 图2 乙 ; c.png | 图3 丙 -->\n"
                           "正文。\n")
        t = spec.parse_task(_write(tmp, doc))
    s = t.sections[0]
    assert s.figures == ["a.png", "b.png", "c.png"], s.figures
    assert s.captions == ["图1 甲", "图2 乙", "图3 丙"], s.captions
    assert [p.name for p in s.figure_paths] == ["a.png", "b.png", "c.png"]
    assert s.has_figures


def test_comma_separated_figures_rejected():
    """逗号列多图是旧写法（会共用一条图注），必须报错并提示改用分号。"""
    with tempfile.TemporaryDirectory() as tmp:
        (Path(tmp) / "assets").mkdir(parents=True, exist_ok=True)
        for n in ("a.png", "b.png"):
            (Path(tmp) / "assets" / n).write_bytes(b"\x89PNG\r\n\x1a\n")
        doc = _minimal_doc("## 甲\n<!-- figure: a.png, b.png | 图1  两图 -->\n正文。\n")
        try:
            spec.parse_task(_write(tmp, doc))
        except ValueError as e:
            assert "分号" in str(e), str(e)
        else:
            raise AssertionError("逗号列多图应当报错")


def test_multi_figure_without_caption_rejected():
    """多图必须逐张写注 —— 少写一条会静默少一条图注，所以在解析层就拦。"""
    with tempfile.TemporaryDirectory() as tmp:
        (Path(tmp) / "assets").mkdir(parents=True, exist_ok=True)
        for n in ("a.png", "b.png"):
            (Path(tmp) / "assets" / n).write_bytes(b"\x89PNG\r\n\x1a\n")
        doc = _minimal_doc("## 甲\n<!-- figure: a.png | 图1 甲 ; b.png -->\n正文。\n")
        try:
            spec.parse_task(_write(tmp, doc))
        except ValueError as e:
            assert "b.png" in str(e) and "图注" in str(e), str(e)
        else:
            raise AssertionError("多图漏写图注应当报错")


def test_caption_directive_lists_each_caption():
    """``caption:`` 也能逐张写（分号分隔），条数必须与图数一致。"""
    with tempfile.TemporaryDirectory() as tmp:
        (Path(tmp) / "assets").mkdir(parents=True, exist_ok=True)
        for n in ("a.png", "b.png"):
            (Path(tmp) / "assets" / n).write_bytes(b"\x89PNG\r\n\x1a\n")
        doc = _minimal_doc("## 甲\n<!-- figure: a.png ; b.png -->\n"
                           "<!-- caption: 图1 甲 ; 图2 乙 -->\n正文。\n")
        t = spec.parse_task(_write(tmp, doc))
        assert t.sections[0].captions == ["图1 甲", "图2 乙"]

        doc = _minimal_doc("## 甲\n<!-- figure: a.png ; b.png -->\n"
                           "<!-- caption: 只有一条 -->\n正文。\n")
        try:
            spec.parse_task(_write(tmp, doc))
        except ValueError as e:
            assert "一张图一条图注" in str(e), str(e)
        else:
            raise AssertionError("图注条数与图数不符应当报错")


def test_missing_second_figure_is_reported():
    """多图里缺任意一张都要报错，并指出是哪一张。"""
    with tempfile.TemporaryDirectory() as tmp:
        (Path(tmp) / "assets").mkdir(parents=True, exist_ok=True)
        (Path(tmp) / "assets" / "a.png").write_bytes(b"\x89PNG\r\n\x1a\n")
        doc = _minimal_doc("## 甲\n"
                           "<!-- figure: a.png | 图1 甲 ; 漏了.png | 图2 乙 -->\n正文。\n")
        try:
            spec.parse_task(_write(tmp, doc))
        except FileNotFoundError as e:
            assert "漏了.png" in str(e)
        else:
            raise AssertionError("缺第二张图应当报错")


def test_too_many_figures_rejected():
    with tempfile.TemporaryDirectory() as tmp:
        names = " ; ".join("f%d.png | 图%d 图" % (i, i)
                           for i in range(spec.MAX_FIGURES + 1))
        doc = _minimal_doc("## 甲\n<!-- figure: %s -->\n正文。\n" % names)
        try:
            spec.parse_task(_write(tmp, doc))
        except ValueError as e:
            assert "最多" in str(e)
        else:
            raise AssertionError("超过单块上限应当报错")


def test_mixed_none_token_rejected():
    """figure 里混进 none 是笔误，必须报错而不是静默少一张图。"""
    with tempfile.TemporaryDirectory() as tmp:
        (Path(tmp) / "assets").mkdir(parents=True, exist_ok=True)
        (Path(tmp) / "assets" / "a.png").write_bytes(b"\x89PNG\r\n\x1a\n")
        doc = _minimal_doc("## 甲\n<!-- figure: a.png | 图1 甲 ; none -->\n正文。\n")
        try:
            spec.parse_task(_write(tmp, doc))
        except ValueError as e:
            assert "none" in str(e)
        else:
            raise AssertionError("混用 none 应当报错")


def test_repeated_figure_directive_rejected():
    """一个小节写两行 figure 是笔误（后一行会把前一行顶掉），必须报错。"""
    with tempfile.TemporaryDirectory() as tmp:
        (Path(tmp) / "assets").mkdir(parents=True, exist_ok=True)
        (Path(tmp) / "assets" / "a.png").write_bytes(b"\x89PNG\r\n\x1a\n")
        doc = _minimal_doc(
            "## 甲\n<!-- figure: a.png -->\n<!-- figure: a.png -->\n正文。\n")
        try:
            spec.parse_task(_write(tmp, doc))
        except ValueError as e:
            assert "多次" in str(e)
        else:
            raise AssertionError("重复 figure 指令应当报错")


def test_figure_none_is_text_only():
    with tempfile.TemporaryDirectory() as tmp:
        doc = _minimal_doc("## 甲\n<!-- figure: none -->\n正文。\n")
        t = spec.parse_task(_write(tmp, doc))
    assert t.sections[0].figures == []
    assert t.sections[0].figure_paths == []
    assert t.sections[0].mode == "text", "无配图必须落到纯文本模式"


def test_mode_directive():
    with tempfile.TemporaryDirectory() as tmp:
        t = _with_fig(tmp, "<!-- figure: img.png -->\n<!-- mode: side -->")
        assert t.sections[0].mode == "side"
        assert t.sections[0].captions == ["甲"], "没写图注时应退回小节标题"


def test_bad_mode_rejected():
    with tempfile.TemporaryDirectory() as tmp:
        try:
            _with_fig(tmp, "<!-- mode: 对角线 -->")
        except ValueError as e:
            assert "mode" in str(e)
        else:
            raise AssertionError("非法 mode 应当报错")


def test_fig_mode_without_figure_rejected():
    with tempfile.TemporaryDirectory() as tmp:
        doc = _minimal_doc("## 甲\n<!-- figure: none -->\n<!-- mode: side -->\n正文。\n")
        try:
            spec.parse_task(_write(tmp, doc))
        except ValueError as e:
            assert "没有配图" in str(e)
        else:
            raise AssertionError("没有配图却要求左文右图，应当报错")


def test_unknown_directive_rejected():
    with tempfile.TemporaryDirectory() as tmp:
        doc = _minimal_doc("## 甲\n<!-- figuer: 拼错了.png -->\n正文。\n")
        try:
            spec.parse_task(_write(tmp, doc))
        except ValueError as e:
            assert "未知指令" in str(e)
        else:
            raise AssertionError("拼错的指令应当报错而不是被静默忽略")


def test_missing_figure_file_is_reported():
    with tempfile.TemporaryDirectory() as tmp:
        doc = _minimal_doc("## 甲\n<!-- figure: 不存在.png -->\n正文。\n")
        try:
            spec.parse_task(_write(tmp, doc))
        except FileNotFoundError as e:
            assert "不存在.png" in str(e) and "配图目录" in str(e)
        else:
            raise AssertionError("配图缺失应当报错")


# ---------------------------------------------------------------- 网格
def test_default_grid_rhythm():
    assert spec.auto_rows(1) == [[0]]
    assert spec.auto_rows(2) == [[0], [1]]
    assert spec.auto_rows(5) == [[0], [1, 2], [3, 4]]
    assert spec.auto_rows(6) == [[0], [1, 2], [3, 4], [5]]


def test_grid_must_cover_every_section():
    doc = _minimal_doc("## 甲\n甲。\n\n## 乙\n乙。\n", extra="grid:\n  - [0]\n")
    with tempfile.TemporaryDirectory() as tmp:
        try:
            spec.parse_task(_write(tmp, doc))
        except ValueError as e:
            assert "不重不漏" in str(e)
        else:
            raise AssertionError("网格漏掉小节应当报错")


def test_row_fracs_must_sum_to_one():
    doc = _minimal_doc("## 甲\n甲。\n\n## 乙\n乙。\n",
                       extra="grid:\n  - [0]\n  - [1]\n"
                             "row_fracs:\n  - [1.0]\n  - [0.6]\n")
    with tempfile.TemporaryDirectory() as tmp:
        try:
            spec.parse_task(_write(tmp, doc))
        except ValueError as e:
            assert "合计" in str(e)
        else:
            raise AssertionError("比例不合 1 应当报错")


def test_row_fracs_length_must_match_grid():
    doc = _minimal_doc("## 甲\n甲。\n\n## 乙\n乙。\n",
                       extra="grid:\n  - [0]\n  - [1]\n"
                             "row_fracs:\n  - [1.0, 1.0]\n  - [1.0]\n")
    with tempfile.TemporaryDirectory() as tmp:
        try:
            spec.parse_task(_write(tmp, doc))
        except ValueError as e:
            assert "row_fracs" in str(e)
        else:
            raise AssertionError("比例个数对不上应当报错")


def test_row_fracs_default_is_equal_split():
    doc = _minimal_doc("## 甲\n甲。\n\n## 乙\n乙。\n\n## 丙\n丙。\n")
    with tempfile.TemporaryDirectory() as tmp:
        t = spec.parse_task(_write(tmp, doc))
    assert t.rows == [[0], [1, 2]]
    assert t.row_fracs == [[1.0], [0.5, 0.5]]


# ---------------------------------------------------------------- 风格覆盖
def test_unknown_style_key_rejected():
    doc = _minimal_doc("## 甲\n甲。\n", extra="style:\n  RADUIS: 0.1\n")
    with tempfile.TemporaryDirectory() as tmp:
        t = spec.parse_task(_write(tmp, doc))
        try:
            C.configure(t)
        except KeyError as e:
            assert "RADUIS" in str(e)
        else:
            raise AssertionError("拼错的风格键应当报错")
        finally:
            C.reset()


def test_style_override_applies_and_resets():
    doc = _minimal_doc("## 甲\n甲。\n",
                       extra="style:\n  TITLE_SZ: 6.5\n  RED: [0x00, 0x80, 0x00]\n")
    with tempfile.TemporaryDirectory() as tmp:
        t = spec.parse_task(_write(tmp, doc))
        default_sz, default_red = C.TITLE_SZ, C.RED
        C.configure(t)
        try:
            assert C.TITLE_SZ == 6.5
            assert C.RED == (0x00, 0x80, 0x00)
            assert C.CHIP_H > default_sz * C.PT2CM, "派生量应随 TITLE_SZ 重算"
        finally:
            C.reset()
        assert C.TITLE_SZ == default_sz and C.RED == default_red


# ---------------------------------------------------------------- 任务定位
def test_find_task_accepts_id_dir_or_file():
    tid = C.TASKS / TASK
    a = spec.find_task(TASK)
    b = spec.find_task(tid)
    c = spec.find_task(tid / "task.md")
    assert a == b == c == (tid / "task.md").resolve()


def test_real_task_parses():
    """随仓库发布的示例任务必须能解析成结构合理的数据。"""
    t = spec.parse_task(C.TASKS / TASK / "task.md")
    assert t.sections, "至少要有一个小节"
    flat = [r for row in t.rows for r in row]
    assert sorted(flat) == list(range(len(t.sections)))
    assert t.char_count() > 0
    assert t.output.endswith(".pptx")
    assert t.template.is_file()
    for sec in t.sections:
        assert len(sec.figures) == len(sec.figure_paths) == len(sec.captions), sec.title
        for name, path in zip(sec.figures, sec.figure_paths):
            assert path.is_file(), name
        if not sec.figures:
            assert sec.mode == "text", "无配图的小节必须是纯文本块"
        else:
            assert all(c.strip() for c in sec.captions), "有配图就必须逐张写图注"


# ---------------------------------------------------------------- 端到端
def test_minimal_task_builds_end_to_end():
    """一份最小的、完全无配图的任务文档也要能跑通全流程。

    这是"通用引擎"的底线：不带任何素材也能出一份 pptx。
    """
    from poster.cli import build
    from pptx import Presentation

    doc = _minimal_doc(
        "## 甲小节\n"
        "这是一段很短的正文，用来验证纯文本块也能排版。**关键数据**标红。\n\n"
        "## 乙小节\n"
        "第二段正文，同样没有配图。\n")
    with tempfile.TemporaryDirectory() as tmp:
        task_path = _write(tmp, doc)
        try:
            res = build(task_ref=task_path, with_preview=False, verbose=False)
            assert res["pptx"].is_file(), "没产出 pptx"
            prs = Presentation(str(res["pptx"]))
            names = [sh.text_frame.text for sh in prs.slides[0].shapes
                     if sh.has_text_frame]
            joined = "".join(names)
            assert "甲小节" in joined and "乙小节" in joined
            assert "关键数据" in joined
        finally:
            C.reset()
