# -*- coding: utf-8 -*-
"""探针 3：用 PowerPoint 自己报的几何量，量出**行距**与**字宽**的真实值。

为什么要这么做
--------------
渲染器不是照 XML 原样画的：`spcPts` 和 `sz` 都可能被量化。而"排进去的字会不会
溢到图上"完全取决于这两个真实值，所以必须量，不能猜。

技巧
----
用 `<a:br/>` **强制**折行数，于是行数是我们说了算的；再由 COM 读回
`TextRange.BoundHeight`。因为 BoundHeight = (n-1) x 行距 + 行框高，
所以两个不同行数的测量值**相减**就能把行框高抵消掉，直接得到行距：
       行距 = (H(n2) - H(n1)) / (n2 - n1)
字宽同理：一整行 GLYPHS 个全角字，BoundWidth / GLYPHS 就是每字推进宽。
（`tools/read_probe.ps1` 里的除数必须与本文件的 GLYPHS 保持一致。）

用法：.venv/Scripts/python.exe tools/probe_metrics.py <输出.pptx>
      （随后 tools/read_probe.ps1 读回几何量）
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'

CJK = "实验表明未浸渍的原木在激光辐照下"     # 16 个全角字，恰好够量字宽
GLYPHS = 16                  # 每行字数；read_probe.ps1 用同一个数做除数

LNVARS: list[tuple[str, str | None]] = [
    ("none", None),
    ("pts400", '<a:lnSpc %s><a:spcPts val="400"/></a:lnSpc>' % NS),
    ("pts500", '<a:lnSpc %s><a:spcPts val="500"/></a:lnSpc>' % NS),
    ("pts540", '<a:lnSpc %s><a:spcPts val="540"/></a:lnSpc>' % NS),
    ("pts550", '<a:lnSpc %s><a:spcPts val="550"/></a:lnSpc>' % NS),
    ("pts552", '<a:lnSpc %s><a:spcPts val="552"/></a:lnSpc>' % NS),
    ("pts560", '<a:lnSpc %s><a:spcPts val="560"/></a:lnSpc>' % NS),
    ("pts600", '<a:lnSpc %s><a:spcPts val="600"/></a:lnSpc>' % NS),
    ("pts1000", '<a:lnSpc %s><a:spcPts val="1000"/></a:lnSpc>' % NS),
    ("pct100", '<a:lnSpc %s><a:spcPct val="100000"/></a:lnSpc>' % NS),
    ("pct133", '<a:lnSpc %s><a:spcPct val="133000"/></a:lnSpc>' % NS),
    ("pct140", '<a:lnSpc %s><a:spcPct val="140000"/></a:lnSpc>' % NS),
    ("pct150", '<a:lnSpc %s><a:spcPct val="150000"/></a:lnSpc>' % NS),
]

N1, N2 = 2, 6                 # 两个测量行数
SIZES = [300, 340, 345, 350, 360]
SZ_BASE = 345


def add_box(slide, x, y, w, h, name, size_pt: float, lines: int,
            lnspc: str | None = None, per_line: int = GLYPHS):
    """放一个文本框；lines 用 <a:br/> 强制。"""
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    bp = tf._txBody.find(qn("a:bodyPr"))
    for k in ("lIns", "rIns", "tIns", "bIns"):
        bp.set(k, "0")
    bp.set("anchor", "t")
    for t in ("a:normAutofit", "a:spAutoFit"):
        e = bp.find(qn(t))
        if e is not None:
            bp.remove(e)
    bp.append(parse_xml('<a:noAutofit %s/>' % NS))

    p = tf.paragraphs[0]
    pPr = p._p.get_or_add_pPr()
    pPr.set("algn", "l")
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    if lnspc:
        pPr.insert(0, parse_xml(lnspc))
    pPr.append(parse_xml('<a:spcBef %s><a:spcPts val="0"/></a:spcBef>' % NS))
    pPr.append(parse_xml('<a:spcAft %s><a:spcPts val="0"/></a:spcAft>' % NS))
    pPr.append(parse_xml('<a:buNone %s/>' % NS))

    txt = CJK[:per_line]
    for i in range(lines):
        if i:
            p._p.append(parse_xml('<a:br %s/>' % NS))
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size_pt)
        r.font.name = "微软雅黑"
        r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        rPr = r._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            rPr.append(parse_xml('<%s %s typeface="微软雅黑"/>' % (tag, NS)))
    return tb


def main(out: Path) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(8.0), Cm(56.0)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    y = 0.30
    for label, lnspc in LNVARS:                       # 每个变体 2 个框
        add_box(slide, 0.30, y, 6.0, 3.5, "A_%s_n%d" % (label, N1),
                SZ_BASE / 100.0, N1, lnspc)
        add_box(slide, 0.30, y + 4.0, 6.0, 3.5, "A_%s_n%d" % (label, N2),
                SZ_BASE / 100.0, N2, lnspc)
        y += 8.0
    for sz in SIZES:                                  # 字宽：一行 GLYPHS 个全角字
        add_box(slide, 0.30, y, 7.0, 0.9, "B_%d" % sz, sz / 100.0, 1,
                None, per_line=GLYPHS)
        y += 1.2

    prs.save(str(out))
    print("已写出：%s（x N1=%d / N2=%d）" % (out, N1, N2))
    for label, _ in LNVARS:
        print("  A_%s" % label)
    for sz in SIZES:
        print("  B_%d = %.2fpt" % (sz, sz / 100.0))


if __name__ == "__main__":
    main(Path(sys.argv[1] if len(sys.argv) > 1 else "tools/_probe/probe_metrics.pptx"))
