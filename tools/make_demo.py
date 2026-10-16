# -*- coding: utf-8 -*-
"""生成示例任务 ``tasks/demo`` 的模板与占位配图。

为什么需要生成器而不是手工塞文件
--------------------------------
示例任务是随仓库发布的**演示素材**，它必须由代码生成：

1. 仓库只发布引擎，不含任何真实赛事材料（见 ``.gitignore``）；
2. 合成模板必须满足 ``template.is_placeholder()`` 的判定条件，否则 ``area: auto``
   探测不到内容区 —— 这是靠代码保证的，不靠手抄坐标；
3. 想换一套演示素材，重跑本脚本即可。

用法::

    .venv/Scripts/python.exe tools/make_demo.py

产出（会覆盖同名文件）::

    tasks/demo/assets/template.pptx      8x20cm 合成模板（含「作品简介」占位框）
    tasks/demo/assets/fig_flow_*.png     4:1  流程条
    tasks/demo/assets/fig_sample_*.png   1.5:1 样品阵列
    tasks/demo/assets/fig_map_*.png      1:1  二维分布
    tasks/demo/assets/fig_bars_*.png     2:1  分组柱状
"""
from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "tasks" / "demo" / "assets"

PX_CM = 200.0                       # 预览/底图渲染分辨率
CANVAS = (8.0, 20.0)                # 合成模板画布（cm）
FONT_FAMILY = "微软雅黑"

# 占位框：必须与 config.PH_TOP（8.52）对齐、宽 >= PH_MIN_W（5.5），
# 这样 area: auto 才能探到内容区 0.74/8.74 -> 7.30/18.30（= config.FALLBACK_AREA）。
PH = (0.69, 8.52, 6.61, 10.20)      # left, top, width, height（cm）

INK = (0x1A, 0x1A, 0x1A)
BLUE = (0x1F, 0x7F, 0xC0)
LBLUE = (0x7F, 0xB5, 0xD9)
PALE = (0xE8, 0xF4, 0xFD)
GRAY = (0x8A, 0x8A, 0x8A)
FONT_PATH = "C:/Windows/Fonts/msyh.ttc"
FONT_BOLD = "C:/Windows/Fonts/msyhbd.ttc"


def font(px: float, bold: bool = False) -> ImageFont.FreeTypeFont:
    for p in ((FONT_BOLD if bold else FONT_PATH), FONT_PATH):
        try:
            return ImageFont.truetype(p, int(px))
        except OSError:
            continue
    return ImageFont.load_default()


def text(d: ImageDraw.ImageDraw, xy, s, f, fill=INK, anchor="la") -> None:
    """先画一层白描边再画本色 —— 免得深色底上的字看不清。"""
    x, y = xy
    for dx, dy in ((-1, 0), (1, 0), (0, -1), (0, 1)):
        d.text((x + dx, y + dy), s, font=f, fill=(255, 255, 255), anchor=anchor)
    d.text((x, y), s, font=f, fill=fill, anchor=anchor)


# ------------------------------------------------------------------ 模板底图
def make_background() -> Image.Image:
    W, H = int(CANVAS[0] * PX_CM), int(CANVAS[1] * PX_CM)
    im = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(im)
    cm = lambda v: v * PX_CM                                   # noqa: E731

    # 外框（模拟赛事模板的装饰边框）
    d.rounded_rectangle([cm(0.22), cm(0.22), cm(7.78), cm(19.78)],
                        radius=int(cm(0.30)), outline=LBLUE, width=int(cm(0.05)))

    # 顶部标题带
    d.rounded_rectangle([cm(0.45), cm(0.55), cm(7.55), cm(1.85)],
                        radius=int(cm(0.18)), fill=PALE, outline=LBLUE,
                        width=max(1, int(cm(0.025))))
    text(d, (cm(4.0), cm(1.20)), "示例赛事 · 作品墙报模板",
         font(cm(0.62)), fill=BLUE, anchor="mm")
    text(d, (cm(0.60), cm(2.10)), "（本模板由 tools/make_demo.py 合成，仅供演示）",
         font(cm(0.30)), fill=GRAY)

    # 「作品简介」区：浅蓝圆角底 + 标签，与占位框位置一致
    x, y, w, h = PH
    d.rounded_rectangle([cm(x), cm(y), cm(x + w), cm(y + h)],
                        radius=int(cm(0.16)), fill=PALE, outline=LBLUE,
                        width=max(1, int(cm(0.02))))
    d.rounded_rectangle([cm(x), cm(y - 0.42), cm(x + 2.10), cm(y - 0.06)],
                        radius=int(cm(0.16)), fill=BLUE)
    text(d, (cm(x + 1.05), cm(y - 0.24)), "作品简介：", font(cm(0.36), True),
         fill=(255, 255, 255), anchor="mm")

    # 底部说明
    text(d, (cm(4.0), cm(19.20)), "示例模板 · 8 x 20 cm（1:10 缩放稿，实印 80 x 200 cm）",
         font(cm(0.28)), fill=GRAY, anchor="mm")
    return im


def make_template(dst: Path) -> None:
    bg = make_background()
    bg_path = dst / "_background.png"
    bg.save(bg_path)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(CANVAS[0]), Cm(CANVAS[1])
    slide = prs.slides.add_slide(prs.slide_layouts[6])          # 空白版式

    # 整页底图：预览取 ppt/media 里最大的一张当背景
    slide.shapes.add_picture(str(bg_path), 0, 0,
                             Cm(CANVAS[0]), Cm(CANVAS[1]))

    # 四个「不许动」的字段（演示引擎只改作品简介区）
    # 显式给字号、并把框放宽到能容下一行：模板字段在 8x20cm 的缩放稿上本来就该是
    # 10~14pt 量级；留空会继承版式默认的 18pt，字比框还高。
    fields = [("作品名称", "示例作品名称", 2.95, 0.85, 14),
              ("项目成员", "张三、李四、王五", 4.00, 0.70, 10),
              ("本科生组", "本科生组", 4.85, 0.70, 10),
              ("作品编号", "DEMO-0001", 5.70, 0.70, 10)]
    for label, val, t, h, sz in fields:
        tb = slide.shapes.add_textbox(Cm(0.60), Cm(t), Cm(6.80), Cm(h))
        tb.text_frame.text = "%s：%s" % (label, val)
        run = tb.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(sz)
        run.font.name = FONT_FAMILY

    # 「作品简介」占位框：故意留一段旧文字，构建时必须被清空
    x, y, w, h = PH
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tb.text_frame.text = ("这里是模板自带的旧文字，构建时会被引擎清空。"
                          "如果没清空，就会和新排版的内容叠印在一起。")
    tb.name = "placeholder_intro"

    prs.save(str(dst / "template.pptx"))
    bg_path.unlink()                                            # 中间产物不入库
    print("写出 %s" % (dst / "template.pptx"))


# ------------------------------------------------------------------ 占位配图
def _axes(d, box, xlabel: str, ylabel: str) -> None:
    """画一个带坐标轴与刻度感的框，让占位图看起来像真的图表。"""
    x0, y0, x1, y1 = box
    d.rectangle([x0, y0, x1, y1], outline=(0x60, 0x60, 0x60), width=2)
    f = font(20)
    text(d, ((x0 + x1) / 2, y1 + 14), xlabel, f, fill=(0x50, 0x50, 0x50), anchor="ma")
    text(d, (x0 - 10, (y0 + y1) / 2), ylabel, f, fill=(0x50, 0x50, 0x50), anchor="rm")


def fig_flow(path: Path) -> None:
    """4:1 的超宽流程条 —— 用来演示「超宽图必须通栏」。"""
    im = Image.new("RGB", (1600, 400), (255, 255, 255))
    d = ImageDraw.Draw(im)
    steps = ["1 预处理", "2 涂布", "3 固化", "4 测试"]
    bw, gap = 300, 100
    x = 60
    for i, s in enumerate(steps):
        d.rounded_rectangle([x, 120, x + bw, 280], radius=16,
                            fill=PALE, outline=BLUE, width=3)
        text(d, (x + bw / 2, 200), s, font(38, True), fill=BLUE, anchor="mm")
        if i < len(steps) - 1:
            d.line([x + bw + 12, 200, x + bw + gap - 12, 200], fill=BLUE, width=4)
            d.polygon([(x + bw + gap - 12, 200), (x + bw + gap - 30, 190),
                       (x + bw + gap - 30, 210)], fill=BLUE)
        x += bw + gap
    text(d, (800, 350), "工艺流程（占位示意图）", font(26), fill=GRAY, anchor="mm")
    im.save(path)


def fig_sample(path: Path) -> None:
    """1.5:1 的样品阵列 —— 演示图条里的第二张。"""
    im = Image.new("RGB", (600, 400), (255, 255, 255))
    d = ImageDraw.Draw(im)
    for r in range(3):
        for c in range(4):
            x, y = 40 + c * 135, 60 + r * 105
            shade = 0xB0 + (r * 4 + c) * 8
            d.rounded_rectangle([x, y, x + 105, y + 75], radius=10,
                                fill=(shade, shade + 20, 235), outline=BLUE, width=2)
    text(d, (300, 380), "样品阵列（占位）", font(24), fill=GRAY, anchor="mm")
    im.save(path)


def fig_map(path: Path) -> None:
    """1:1 的二维分布 —— 演示「方图走文左图右」。"""
    im = Image.new("RGB", (800, 800), (255, 255, 255))
    d = ImageDraw.Draw(im)
    box = (110, 70, 740, 640)
    _axes(d, box, "x / a.u.", "y / a.u.")
    x0, y0, x1, y1 = box
    n = 10
    for i in range(n):
        for j in range(n):
            v = (i * 7 + j * 13) % 10 / 9.0            # 固定的伪数据，不要随机
            shade = int(0xE0 - 0x90 * v)
            d.rectangle([x0 + i * (x1 - x0) / n, y0 + j * (y1 - y0) / n,
                         x0 + (i + 1) * (x1 - x0) / n, y0 + (j + 1) * (y1 - y0) / n],
                        fill=(shade, shade, 0xF0), outline=(255, 255, 255))
    text(d, (400, 720), "二维分布（占位）", font(26), fill=GRAY, anchor="mm")
    im.save(path)


def fig_bars(path: Path) -> None:
    """2:1 的分组柱状 —— 演示「中等宽高比按块宽定排法」。"""
    im = Image.new("RGB", (1200, 600), (255, 255, 255))
    d = ImageDraw.Draw(im)
    box = (120, 70, 1120, 500)
    _axes(d, box, "组别", "响应 / a.u.")
    x0, y0, x1, y1 = box
    vals = [0.35, 0.62, 0.48, 0.81, 0.29, 0.55]
    bw = (x1 - x0) / (len(vals) * 1.6)
    for i, v in enumerate(vals):
        bx = x0 + (x1 - x0) * (i + 0.3) / len(vals)
        by = y1 - (y1 - y0) * v
        d.rectangle([bx, by, bx + bw, y1], fill=BLUE)
        text(d, (bx + bw / 2, by - 22), "%.2f" % v, font(20),
             fill=(0x40, 0x40, 0x40), anchor="mm")
    text(d, (620, 560), "分组响应（占位数据）", font(26), fill=GRAY, anchor="mm")
    im.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    make_template(OUT)
    figs = [("fig_flow_1600x400.png", fig_flow),
            ("fig_sample_600x400.png", fig_sample),
            ("fig_map_800x800.png", fig_map),
            ("fig_bars_1200x600.png", fig_bars)]
    for name, fn in figs:
        p = OUT / name
        fn(p)
        print("写出 %s（%dx%d，%.0f KB）"
              % (p.relative_to(ROOT), *Image.open(p).size, p.stat().st_size / 1024))


if __name__ == "__main__":
    main()
