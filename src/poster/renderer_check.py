# -*- coding: utf-8 -*-
"""渲染器真值复核：让 PowerPoint 自己说「排进去的字有没有压到图上」。

为什么必须有
------------
本引擎只能用 Pillow 量**字宽**、按模型推**行距**；而渲染器有自己的一套
（断行策略、行距量化、字宽推进）。两者一旦有差，就会出现"末行被下面的配图
盖掉"这种只有真渲染器才看得见的事故 —— 2026-10-14 就是这样翻车的：
PowerPoint 把 ``spcPts 552``（5.52pt）按 6.0pt 排，一行多 8.7% 高，整块多一行，
第 10 行整行钻到它下面那张图底下。**模型里一切自洽，只有渲染器知道不对劲。**

做法
----
借 COM 让 PowerPoint 打开产物、把每个形状的几何量导成数字 TSV
（``tools/ppt_dump.ps1``；其中的 ``Bound*`` 是**渲染器自己算出来的文字墨迹框**），
再判读三件事：
  1. 每个文本框的墨迹底有没有落在下面那张图的范围内；
  2. 墨迹高有没有超过它自己声明的框高（超了就是"迟早会溢出去"）；
  3. 整份文件能不能被 PowerPoint 打开（打不开本身就是硬伤）。

没有装 PowerPoint 时 :func:`available` 返回 False，调用方应把结果记为「跳过」。
"""
from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

PT_CM = 2.54 / 72.0
REPO = Path(__file__).resolve().parents[2]
DUMP_PS1 = REPO / "tools" / "ppt_dump.ps1"


class ShapeGeom:
    """PowerPoint 报回的一个形状的几何量（已换算成 cm）。"""

    __slots__ = ("idx", "kind", "has_tf", "x", "y", "w", "h",
                 "bx", "by", "bw", "bh", "tlen", "npar", "name")

    def __init__(self, row: list[str]):
        self.idx = int(row[0])
        self.kind = row[1]
        self.has_tf = row[2] == "1"
        self.x = float(row[3]) * PT_CM
        self.y = float(row[4]) * PT_CM
        self.w = float(row[5]) * PT_CM
        self.h = float(row[6]) * PT_CM
        self.bx = float(row[7]) * PT_CM
        self.by = float(row[8]) * PT_CM
        self.bw = float(row[9]) * PT_CM
        self.bh = float(row[10]) * PT_CM
        self.tlen = int(float(row[11]))
        self.npar = int(float(row[12]))
        self.name = row[13] if len(row) > 13 else ""

    @property
    def b_bottom(self) -> float:
        return self.by + self.bh

    @property
    def b_right(self) -> float:
        return self.bx + self.bw

    def __repr__(self) -> str:
        return "#%d %s(%.2f,%.2f %.2fx%.2f)" % (self.idx, self.kind, self.x, self.y, self.w, self.h)


def available() -> bool:
    """本机能不能做渲染器复核（要 Windows + PowerPoint + PowerShell）。"""
    if sys.platform != "win32" or not DUMP_PS1.is_file():
        return False
    return shutil.which("powershell") is not None


def dump_shapes(pptx) -> list[ShapeGeom]:
    """让 PowerPoint 打开 pptx 并导出所有形状几何量。

    **打开的是产物的副本，不是产物本身。** PowerPoint 会给它打开的文件加共享锁、
    并留下 ``~$<名>.pptx`` owner 文件；只要实例没退干净，**下一轮构建就写不进
    产物**（``pptx`` 打包时会抛 ``PermissionError``，F3 复现性跟着红）。
    复核只读内容，副本完全等价，于是产物再也不可能被渲染器锁住。
    """
    out = Path(tempfile.gettempdir()) / "poster_ppt_dump.tsv"
    if out.exists():
        out.unlink()
    probe_dir = Path(tempfile.mkdtemp(prefix="poster_probe_"))
    probe = probe_dir / Path(pptx).name
    shutil.copy2(pptx, probe)
    try:
        cmd = ["powershell", "-NoProfile", "-File", str(DUMP_PS1),
               str(probe), str(out)]
        # PowerShell 的 stdout 是本地代码页（中文 GBK），不能按 UTF-8 解；只取字节再容错解码
        r = subprocess.run(cmd, capture_output=True, timeout=300)
        so = ((r.stdout or b"") + (r.stderr or b"")).decode("utf-8", "replace")
        if r.returncode != 0 or not out.exists():
            raise RuntimeError("PowerPoint 打不开或导出失败（exit=%s）：\n%s" % (r.returncode, so))
    finally:
        shutil.rmtree(probe_dir, ignore_errors=True)
    rows = [ln.split("\t") for ln in out.read_text(encoding="utf-8-sig").splitlines() if ln.strip()]
    return [ShapeGeom(x) for x in rows]


def verify(pptx, tol: float = 0.02, n_tpl: int = 0) -> list[str]:
    """返回「渲染器视角下的问题」列表；空列表 = 通过。

    ``n_tpl`` = 模板原有形状数。python-pptx 是**只追加**的，所以产物里前 ``n_tpl``
    个形状原封不动来自模板 —— 那些框的尺寸是模板作者定的，不归引擎管，
    因此**"溢出声明框"只判引擎排进去的文字**（与 compliance 的 ``ours_only`` 同一口径）。
    但"文字压到图"对**全部**文字生效：万一我们排的图盖住了模板自带的字段，那是真事故。
    """
    shapes = dump_shapes(pptx)
    canvas = max((s.w * s.h for s in shapes), default=1.0)
    # 整张底图不是"配图"；空文本框会报出空段落的行框（不是真墨迹），都要滤掉
    pics = [s for s in shapes if s.kind == "PICTURE" and s.w * s.h < 0.9 * canvas]
    problems: list[str] = []

    for i, t in enumerate(shapes):
        if not (t.has_tf and t.tlen > 0 and t.bh > 0):
            continue
        if i >= n_tpl and t.bh > t.h + tol:
            problems.append(
                "文字溢出声明框：%s y=%.3f 声明高 %.3f < 墨迹高 %.3f（超 %.3fcm）"
                % (t, t.y, t.h, t.bh, t.bh - t.h))
        for p in pics:
            if p.y >= t.b_bottom - tol or p.y + p.h <= t.by:
                continue
            if p.x >= t.b_right - tol or p.x + p.w <= t.bx:
                continue
            problems.append(
                "文字压到图：%s 墨迹 y %.3f..%.3f 与图 x%.2f y%.2f %.2fx%.2f 重叠 %.3fcm"
                % (t, t.by, t.b_bottom, p.x, p.y, p.w, p.h, t.b_bottom - p.y))
    return problems


def template_shape_count(template) -> int:
    """模板原有形状数（= 产物里前几个"不归引擎管"的形状）。"""
    from pptx import Presentation

    return len(Presentation(str(template)).slides[0].shapes)


def summary(pptx, tol: float = 0.02, n_tpl: int = 0) -> str:
    """给人看的复核摘要（供合规报告用）。"""
    problems = verify(pptx, tol, n_tpl)
    if problems:
        return "渲染器复核不通过：%s" % "；".join(problems[:4])
    return "PowerPoint 实排复核通过：无文字压图、无文字溢出声明框"
